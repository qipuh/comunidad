from fastapi import APIRouter
from fastapi.responses import Response

router = APIRouter(prefix="/api", tags=["plantilla"])


@router.get("/plantilla")
def descargar_plantilla():
    """Descargar plantilla de importación en Excel"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        import io

        wb = Workbook()

        # ════════ PESTAÑA 1: USUARIOS ════════
        ws = wb.active
        ws.title = "Usuarios"

        # Estilos
        header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        border = Border(
            left=Side(style='thin', color='D0D0D0'),
            right=Side(style='thin', color='D0D0D0'),
            top=Side(style='thin', color='D0D0D0'),
            bottom=Side(style='thin', color='D0D0D0')
        )

        data_alignment = Alignment(horizontal='left', vertical='center')
        center_alignment = Alignment(horizontal='center', vertical='center')

        # Encabezados
        headers = [
            'num_padron',
            'apellido_paterno',
            'apellido_materno',
            'nombres',
            'sexo',
            'dni',
            'fecha_nacimiento',
            'estado_civil'
        ]

        ws.append(headers)

        # Formatear encabezados
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # Datos de ejemplo
        ejemplos = [
            ['P001', 'García', 'López', 'Juan Carlos', 'M', '12345678', '1990-01-15', 'Soltero'],
            ['P002', 'Pérez', 'Rodríguez', 'María Elena', 'F', '87654321', '1988-06-20', 'Casada'],
            ['P003', 'López', 'Martínez', 'Carlos Alberto', 'M', '11223344', '1995-03-10', 'Soltero'],
            ['P004', 'González', 'Fernández', 'Ana Patricia', 'F', '55667788', '1992-08-25', 'Casada'],
            ['P005', 'Rodríguez', 'García', 'Pedro Luis', 'M', '99887766', '1989-12-05', 'Divorciado'],
        ]

        for ejemplo in ejemplos:
            ws.append(ejemplo)

        # Formatear datos
        for row in range(2, len(ejemplos) + 2):
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = border
                cell.font = Font(name='Calibri', size=11)
                if col == 5:  # Sexo - centrado
                    cell.alignment = center_alignment
                elif col == 6:  # DNI - centrado
                    cell.alignment = center_alignment
                else:
                    cell.alignment = data_alignment

        # Ajustar ancho de columnas
        widths = {
            'A': 12,  # num_padron
            'B': 18,  # apellido_paterno
            'C': 18,  # apellido_materno
            'D': 20,  # nombres
            'E': 8,   # sexo
            'F': 12,  # dni
            'G': 18,  # fecha_nacimiento
            'H': 15   # estado_civil
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.row_dimensions[1].height = 25

        # ════════ PESTAÑA 2: INSTRUCCIONES ════════
        ws_inst = wb.create_sheet('Instrucciones')

        title_font = Font(name='Calibri', size=16, bold=True, color='4F46E5')
        subtitle_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
        subtitle_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        body_font = Font(name='Calibri', size=11)

        # Título
        ws_inst['A1'] = 'GUÍA DE IMPORTACIÓN DE USUARIOS'
        ws_inst['A1'].font = title_font
        ws_inst.row_dimensions[1].height = 20

        # Secciones
        row = 3
        sections = [
            ('CAMPOS OBLIGATORIOS:', [
                '• dni: Número de DNI (exactamente 8 dígitos, solo números)',
                '• nombres: Nombres completos del usuario',
                '• apellido_paterno: Apellido paterno (requerido)',
                '• apellido_materno: Apellido materno (requerido)',
            ]),
            ('CAMPOS OPCIONALES:', [
                '• num_padron: Número de padrón o identificador',
                '• sexo: M, F, Masculino o Femenino',
                '• fecha_nacimiento: Formato DD/MM/YYYY o YYYY-MM-DD',
                '• estado_civil: Soltero, Casado, Divorciado, Viudo',
            ]),
            ('NOTAS IMPORTANTES:', [
                '✓ El DNI debe ser único (sin duplicados en el archivo)',
                '✓ Si un DNI ya existe en el sistema, se actualizará con los nuevos datos',
                '✓ Las contraseñas iniciales serán el DNI (ejemplo: 12345678)',
                '✓ Los usuarios se crearán con estado "Activo"',
                '✓ El rol inicial será "Usuario" (se puede cambiar después)',
            ]),
            ('PASOS PARA IMPORTAR:', [
                '1. Completa los datos en la pestaña "Usuarios"',
                '2. Guarda el archivo',
                '3. En la aplicación: Click en "Importar"',
                '4. Selecciona este archivo',
                '5. Revisa la previsualización (verás errores si los hay)',
                '6. Selecciona los usuarios a importar',
                '7. Click en "Importar" y confirma',
            ]),
        ]

        for section_title, items in sections:
            ws_inst[f'A{row}'] = section_title
            ws_inst[f'A{row}'].font = subtitle_font
            ws_inst[f'A{row}'].fill = subtitle_fill
            ws_inst.row_dimensions[row].height = 18
            row += 1

            for item in items:
                ws_inst[f'A{row}'] = item
                ws_inst[f'A{row}'].font = body_font
                ws_inst[f'A{row}'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                ws_inst.row_dimensions[row].height = 20
                row += 1

            row += 1

        ws_inst.column_dimensions['A'].width = 95

        output = io.BytesIO()
        wb.save(output)

        return Response(
            content=output.getvalue(),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename=plantilla_usuarios.xlsx'}
        )
    except Exception as e:
        return {"error": str(e)}
