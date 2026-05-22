"""
Script para generar un archivo Excel de plantilla para importar usuarios.
Ejecutar como: python generate_excel_template.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

def generar_plantilla():
    """Genera un archivo Excel de plantilla para la importación de usuarios."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Usuarios"

    # Definir estilos
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_alignment = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # Agregar encabezados
    headers = ['num_padron', 'apellido_paterno', 'apellido_materno', 'nombres', 'sexo', 'dni', 'fecha_nacimiento', 'estado_civil']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = thin_border

    # Agregar datos de ejemplo
    ejemplo_datos = [
        ['P001', 'García', 'López', 'Juan', 'M', '12345678', '1990-01-15', 'Soltero'],
        ['P002', 'Pérez', 'Rodríguez', 'María', 'F', '87654321', '1988-06-20', 'Casada'],
        ['P003', 'López', 'Martínez', 'Carlos', 'M', '11223344', '1995-03-10', 'Soltero'],
    ]

    for row_idx, datos in enumerate(ejemplo_datos, 2):
        for col_idx, valor in enumerate(datos, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = valor
            cell.alignment = left_alignment if col_idx >= 1 else center_alignment
            cell.border = thin_border

    # Ajustar ancho de columnas
    anchos = [15, 18, 18, 18, 8, 12, 18, 15]
    for col_idx, ancho in enumerate(anchos, 1):
        ws.column_dimensions[chr(64 + col_idx)].width = ancho

    # Altura de encabezado
    ws.row_dimensions[1].height = 25

    # Agregar hoja de instrucciones
    ws_info = wb.create_sheet("Instrucciones", 0)
    instrucciones = [
        ["GUÍA DE IMPORTACIÓN DE USUARIOS"],
        [""],
        ["Campos Obligatorios:"],
        ["  • dni: Número de DNI (8 dígitos, solo números)"],
        ["  • apellido_paterno: Apellido paterno del usuario"],
        ["  • apellido_materno: Apellido materno del usuario"],
        ["  • nombres: Nombres del usuario"],
        [""],
        ["Campos Opcionales:"],
        ["  • num_padron: Número de padrón (identificador)"],
        ["  • sexo: Sexo (M, F, Masculino o Femenino)"],
        ["  • fecha_nacimiento: Formato DD/MM/YYYY o YYYY-MM-DD"],
        ["  • estado_civil: Estado civil del usuario"],
        [""],
        ["Notas Importantes:"],
        ["  ✓ El DNI debe ser único (no pueden haber duplicados en el archivo)"],
        ["  ✓ Si un DNI ya existe en el sistema:"],
        ["    - Se mostrará en la previsualización como 'Actualizar'"],
        ["    - Se puede optar por actualizar los datos existentes"],
        ["  ✓ Las contraseñas se crearán automáticamente (DNI = contraseña)"],
        ["  ✓ Los usuarios se crearán con estado 'Activo'"],
        ["  ✓ El rol será 'Usuario' (se puede cambiar después manualmente)"],
        [""],
        ["Pasos para importar:"],
        ["  1. Completa los datos en la pestaña 'Usuarios'"],
        ["  2. Sube el archivo desde 'Importar' en la aplicación"],
        ["  3. Revisa la previsualización (verás errores y duplicados)"],
        ["  4. Aplica filtros si es necesario"],
        ["  5. Selecciona los usuarios a importar"],
        ["  6. Confirma la importación"],
    ]

    for row_idx, instruccion in enumerate(instrucciones, 1):
        cell = ws_info.cell(row=row_idx, column=1)
        cell.value = instruccion[0] if instruccion else ""
        if row_idx == 1:
            cell.font = Font(bold=True, size=14, color="4F46E5")
        elif instruccion and instruccion[0].startswith("  "):
            cell.font = Font(size=10)
        else:
            cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    ws_info.column_dimensions['A'].width = 80

    # Guardar archivo
    output_path = Path("uploads/usuarios/plantilla_importacion.xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(f"[SUCCESS] Plantilla de importación creada: {output_path}")

if __name__ == "__main__":
    try:
        generar_plantilla()
    except Exception as e:
        print(f"[ERROR] Error al generar la plantilla: {e}")
